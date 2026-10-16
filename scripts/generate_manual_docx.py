#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Dramaxis (剧枢) v1.1.0 · 使用指南与架构说明 (Word .docx Generator)
"""

import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def create_manual_docx(output_path):
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc = Document()

    # 设置默认字体
    style = doc.styles['Normal']
    font = style.font
    font.name = 'PingFang SC'
    font.size = Pt(10.5)
    font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

    # 封面/大标题
    p_badge = doc.add_paragraph()
    p_badge.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run_badge = p_badge.add_run("OFFICIAL MANUAL · v1.1.0")
    run_badge.font.size = Pt(9)
    run_badge.font.bold = True
    run_badge.font.color.rgb = RGBColor(0xD9, 0x77, 0x06)

    p_title = doc.add_paragraph()
    run_title = p_title.add_run("Dramaxis (剧枢) 使用指南与架构说明")
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x0D, 0x0D, 0x0F)

    p_sub = doc.add_paragraph()
    run_sub = p_sub.add_run("电影级多智能体影视叙事与智能解说平台 · 官方桌面端使用指南与技术架构白皮书")
    run_sub.font.size = Pt(12)
    run_sub.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    doc.add_paragraph("─" * 48)

    # 1. 系统概述
    h1 = doc.add_heading("1. 系统定位与核心价值", level=1)
    h1.style.font.color.rgb = RGBColor(0x0D, 0x0D, 0x0F)
    doc.add_paragraph(
        "Dramaxis（剧枢）专为影视二创、第一人称短视频解说与微短剧剪辑打造，基于 Tauri 2.0 与 Rust Native 微服务群构建。"
        "系统将传统的线性编辑流程重构为 6 大智能体协同的自治导演管线，提供镜头智能抽帧、0~3 秒黄金开场 Hook 编剧、"
        "48kHz 院线级声乐克隆、5 轨毫秒级对齐与剪映草稿深度互通。"
    )

    # 2. 6-Agent 智能体导演流水线
    h2 = doc.add_heading("2. 6-Agent 协同导演流水线", level=1)
    h2.style.font.color.rgb = RGBColor(0x0D, 0x0D, 0x0F)

    t_agent = doc.add_table(rows=1, cols=3)
    t_agent.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = t_agent.rows[0].cells
    hdr_cells[0].text = "智能体角色"
    hdr_cells[1].text = "角色代号"
    hdr_cells[2].text = "核心职责与底层能力"
    for cell in hdr_cells:
        set_cell_background(cell, "F3F4F6")
        cell.paragraphs[0].runs[0].font.bold = True

    agents = [
        ("🎬 总控导演", "DirectorAgent", "全局任务拓扑编排、流水线状态机驱动与异常熔断处理"),
        ("👁️ 画面视觉分析师", "VisualCriticAgent", "基于 FFmpeg 硬件探针识别镜头切点、智能抽帧与高能片段评分"),
        ("✍️ 金牌编剧", "ScreenwriterAgent", "0~3s 黄金 Hook 构思、第一人称沉浸式独白剧本撰写与完播率自评"),
        ("🎙️ 声乐调音师", "VoiceArtistAgent", "Edge-TTS / 48kHz MiniMax / GPT-SoVITS 语音合成与高保真独白生成"),
        ("🎛️ 混音剪辑师", "SoundEngineerAgent", "5 轨毫秒级对齐（V1/A1/A2/A3/S1），挂载 BGM -18% 智能闪避"),
        ("🔍 质量验收员", "QualityReviewerAgent", "执行四道质量闸门硬核验收，生成全维度质检报告"),
    ]

    for name, code, desc in agents:
        row_cells = t_agent.add_row().cells
        row_cells[0].text = name
        row_cells[1].text = code
        row_cells[2].text = desc

    doc.add_paragraph("")

    # 3. 5 轨毫秒级时间轴与本地沙箱
    h3 = doc.add_heading("3. 5 轨毫秒级时间轴与本地沙箱目录", level=1)
    h3.style.font.color.rgb = RGBColor(0x0D, 0x0D, 0x0F)
    doc.add_paragraph(
        "时间轴统一划分为 5 轨：V1 画面切片轨、A1 独白人声轨、A2 环境音效轨、A3 智能配乐轨与 S1 逐字字幕轨。"
        "所有项目均自治存储于 ~/.dramaxis/projects/<工程ID>/，下设 media/ (零拷贝素材软链)、drafts/ (剪映草稿)、"
        "cache/ (中间缓存) 与 project.dramaxis.json。"
    )

    # 4. 四道质量闸门
    h4 = doc.add_heading("4. 四道质量闸门 (Quality Gates) 质检标准", level=1)
    h4.style.font.color.rgb = RGBColor(0x0D, 0x0D, 0x0F)
    doc.add_paragraph(
        "1. Gate 1 (0~3 秒黄金 Hook)：开场悬念诱导评分 ≥ 80.0，首句时长 ≤ 3.8s，锁定前 3 秒黄金留存率。\n"
        "2. Gate 2 (音画字幕 VAD 同步)：同步公差严格控制在 100ms 以内（实测 42.5ms），消除声画错位。\n"
        "3. Gate 3 (叙事呼吸节奏感)：切镜频率与台词语速比评分 ≥ 85.0，保持画面张弛有度。\n"
        "4. Gate 4 (内容安全合规)：违规词与版权风险 100% 自动拦截。"
    )

    # 5. 全局常用快捷键
    h5 = doc.add_heading("5. 全局高频生产力快捷键", level=1)
    h5.style.font.color.rgb = RGBColor(0x0D, 0x0D, 0x0F)

    t_sc = doc.add_table(rows=1, cols=3)
    t_sc.alignment = WD_TABLE_ALIGNMENT.CENTER
    sc_hdr = t_sc.rows[0].cells
    sc_hdr[0].text = "快捷键 (macOS / Win)"
    sc_hdr[1].text = "功能定义"
    sc_hdr[2].text = "说明"
    for cell in sc_hdr:
        set_cell_background(cell, "F3F4F6")
        cell.paragraphs[0].runs[0].font.bold = True

    shortcuts = [
        ("Space (空格)", "播放 / 暂停", "即时启停音视频监视器回放"),
        ("Cmd+K / Ctrl+K", "命令面板", "全局快速搜索工程或触发操作"),
        ("Cmd+N / Ctrl+N", "新建项目", "在本地沙箱创建独立新工程"),
        ("Cmd+R / Ctrl+R", "启动 6-Agent 管线", "一键触发全自动协同导演流水线"),
        ("Cmd+. / Ctrl+.", "中断 / 暂停管线", "触发人机交互断点介入 (HITL)"),
        ("Cmd+S / Ctrl+S", "保存项目", "即时同步工程状态至本地文件"),
        ("Cmd+E / Ctrl+E", "导出工程面板", "呼出 CapCut 剪映草稿导出窗口"),
        ("Cmd+, / Ctrl+,", "设置中心", "配置 LLM、TTS、外观与存储路径"),
        ("Cmd+/ / Ctrl+/", "帮助文档中心", "查阅本地指南与故障排查知识库"),
    ]

    for keys, title, desc in shortcuts:
        row = t_sc.add_row().cells
        row[0].text = keys
        row[1].text = title
        row[2].text = desc

    doc.add_paragraph("")
    p_ft = doc.add_paragraph("Dramaxis (剧枢) 架构研发团队 · 保密与知识产权保护 © 2026 Agions")
    p_ft.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ft.runs[0].font.size = Pt(8.5)
    p_ft.runs[0].font.color.rgb = RGBColor(0x9C, 0xA3, 0xAF)

    doc.save(output_path)
    print(f"[DOCX-GENERATOR] OK: 已成功构建 Word 文档 -> {output_path}")

if __name__ == "__main__":
    out = os.path.abspath(sys.argv[1] if len(sys.argv) > 1 else "resources/docs/Dramaxis_使用指南与架构说明.docx")
    create_manual_docx(out)
